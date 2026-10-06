"""Import the NS order workbook into ctl_* tables.

Reads the spreadsheet once, merges duplicate rows into one order, and prints
a report. With SUPABASE_URL, SUPABASE_ANON_KEY, CONTROLE_EMAIL and
CONTROLE_PASSWORD set, it also replaces the controle tables.
"""

from __future__ import annotations

import json
import math
import os
import re
import unicodedata
import uuid
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(r"c:\Users\Pichau\Downloads\[NS] Acompanhamento de Pedidos.xlsx")
SENTINEL = "0001-01-01"
MONEY_FIELDS = (
    "sale_amount",
    "purchase_amount",
    "payment_fee",
    "shipping_cost",
    "import_tax",
)
TRACKING_RE = re.compile(r"^[A-Z]{2}\d{6,}[A-Z]{2}$", re.I)
MONTHS = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}
SHEET_MONTH = {
    "Financeiro 0826": "2026-08-01",
    "Financeiro 0926": "2026-09-01",
    "Financeiro 1026": "2026-10-01",
}


def fold(value: str) -> str:
    text = unicodedata.normalize("NFKD", value)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).strip().lower()


def parse_money(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return round(float(value), 2)
    text = str(value).strip()
    if not text or text.startswith("=") or text.startswith("#"):
        return None
    text = text.replace("R$", "").replace(" ", "")
    if re.search(r"[A-Za-z]", text):
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return round(float(text), 2)
    except ValueError:
        return None


def parse_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def parse_datetime(value):
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).isoformat(sep=" ")
    return None


def parse_bool(value):
    if isinstance(value, bool):
        return value
    if value is None:
        return None
    text = fold(str(value))
    if text in {"true", "sim", "1", "yes", "verdadeiro"}:
        return True
    if text in {"false", "nao", "0", "no"}:
        return False
    return None


def cell(row, index):
    if index >= len(row):
        return None
    return row[index]


def clean_text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def parse_key(value):
    if value is None or isinstance(value, (datetime, date, bool)):
        return None, None
    if isinstance(value, float) and value.is_integer():
        value = str(int(value))
    text = str(value).strip()
    if not text or text.startswith("="):
        return None, None
    if re.match(r"\d{4}-\d{2}-\d{2}", text):
        return None, None
    folded = fold(text)
    if folded in {"pedido", "data", "qtde", "referencia", "status", "id pedido"}:
        return None, None
    labeled = re.match(r"pedido\s+(\d+)\b(.*)$", folded, re.I)
    original = re.sub(r"\.0$", "", text)
    if labeled:
        hint = labeled.group(2).strip(" -()")
        return labeled.group(1), original
    bare = re.match(r"^(\d+)\b(.*)$", folded)
    if bare and bare.group(1):
        return bare.group(1), original
    slug = re.sub(r"[^a-z0-9]+", "-", folded).strip("-")
    if not slug:
        return None, None
    return f"label:{slug}", original


def hint_from_label(label, key):
    if not label:
        return None
    match = re.match(rf"(?i)^pedido\s+{re.escape(key)}\b\s*(.*)$", label.strip())
    if not match:
        match = re.match(rf"^{re.escape(key)}\s*[\(](.+)[\)]\s*$", label.strip())
        if match:
            return match.group(1).strip() or None
        return None
    hint = match.group(1).strip(" -()")
    return hint or None


def months_in_text(text):
    folded = fold(text)
    found = []
    for name, month in MONTHS.items():
        for match in re.finditer(name + r"\s+de\s+(\d{4})", folded):
            found.append(date(int(match.group(1)), month, 1))
    return found


def choose_month(candidates, previous):
    if not candidates:
        return None
    if previous:
        prev = date.fromisoformat(previous)
        year = prev.year if prev.month > 1 else prev.year - 1
        month = prev.month - 1 or 12
        target = date(year, month, 1)
        if target in candidates:
            return target.isoformat()
        earlier = [item for item in candidates if item < prev]
        if earlier:
            return max(earlier).isoformat()
    return max(candidates).isoformat()


def split_blocks(rows):
    marks = []
    for index, row in enumerate(rows):
        text = " ".join(str(item) for item in row if isinstance(item, str))
        found = months_in_text(text)
        if found and "controle financeiro" in fold(text):
            marks.append((index, found))
    if not marks:
        return [(None, rows)]
    blocks = []
    previous = None
    for position, (index, found) in enumerate(marks):
        month = choose_month(found, previous)
        end = marks[position + 1][0] if position + 1 < len(marks) else len(rows)
        blocks.append((month, rows[index:end]))
        previous = month
    return blocks


def find_meta(rows):
    for index, row in enumerate(rows[:-1]):
        nxt = rows[index + 1]
        for column, value in enumerate(row):
            if not isinstance(value, str) or fold(value) != "meta":
                continue
            if column >= len(nxt):
                continue
            amount = parse_money(nxt[column])
            if amount is not None and amount >= 100000:
                return amount
    return None


def blank_order(key, label, flow, month):
    month_key = month or SENTINEL
    return {
        "id": str(uuid.uuid4()),
        "order_key": key,
        "label": label or key,
        "flow": flow,
        "origin": None,
        "reference": None,
        "product_name": None,
        "commercial_status": None,
        "sourcing_status": None,
        "purchased": False,
        "cpf_linked": False,
        "tax_paid": False,
        "delivered": False,
        "tracking_code": None,
        "tracking_situation": None,
        "tracking_alert": None,
        "tracking_correios": None,
        "tracking_event_at": None,
        "purchase_date": None,
        "payment_date": None,
        "sale_amount": None,
        "purchase_amount": None,
        "payment_fee": None,
        "shipping_cost": None,
        "import_tax": None,
        "supplier_days": None,
        "supplier_ref": None,
        "notes_internal": None,
        "notes_robot": None,
        "notes_human": None,
        "finance_month": month,
        "finance_month_key": month_key,
        "priority": 0,
    }


class Importer:
    def __init__(self):
        self.orders = {}
        self.goals = {}
        self.pending_taxes = []
        self.inventory = []
        self.trade_ins = []
        self.cancellations = []
        self.report = {
            "linhas_lidas": 0,
            "taxas_fundidas": 0,
            "taxas_sem_venda": 0,
            "itens_somados": 0,
            "linhas_sem_numero": 0,
            "linhas_sem_numero_amostra": [],
            "conflitos": 0,
            "conflitos_amostra": [],
            "metas": {},
            "meses": {},
        }

    def identity(self, order):
        return (order["order_key"], order["flow"], order["finance_month_key"])

    def conflict(self, key, field, previous, new, source):
        self.report["conflitos"] += 1
        if len(self.report["conflitos_amostra"]) < 30:
            self.report["conflitos_amostra"].append(
                {
                    "pedido": key,
                    "campo": field,
                    "anterior": previous,
                    "novo": new,
                    "origem": source,
                }
            )

    def assign_money(self, existing, incoming, priority, source):
        for field in MONEY_FIELDS:
            new = incoming.get(field)
            if new is None:
                continue
            old = existing.get(field)
            if old is None:
                existing[field] = new
                continue
            if abs(old - new) <= 0.009:
                continue
            self.conflict(existing["order_key"], field, old, new, source)
            if priority >= existing["priority"]:
                existing[field] = new

    def fill_text(self, existing, incoming, fields, force=False):
        for field in fields:
            new = incoming.get(field)
            if not new:
                continue
            if force or not existing.get(field):
                existing[field] = new

    def upsert_sale(self, incoming, priority, source):
        ident = self.identity(incoming)
        existing = self.orders.get(ident)
        if existing is None:
            incoming["priority"] = priority
            self.orders[ident] = incoming
            return incoming
        new_ref = incoming.get("reference")
        old_ref = existing.get("reference")
        distinct_item = bool(
            new_ref
            and old_ref
            and fold(new_ref) != fold(old_ref)
            and incoming.get("sale_amount") is not None
            and existing.get("sale_amount") is not None
        )
        if distinct_item:
            for field in MONEY_FIELDS:
                if incoming.get(field) is not None:
                    existing[field] = round((existing.get(field) or 0) + incoming[field], 2)
            if new_ref not in old_ref:
                existing["reference"] = f"{old_ref} / {new_ref}"
            self.report["itens_somados"] += 1
            return existing
        self.assign_money(existing, incoming, priority, source)
        if priority >= existing["priority"]:
            if incoming.get("sourcing_status"):
                existing["sourcing_status"] = incoming["sourcing_status"]
            self.fill_text(
                existing,
                incoming,
                ("label", "reference", "product_name", "origin", "tracking_code"),
                force=True,
            )
            existing["priority"] = priority
        else:
            self.fill_text(
                existing,
                incoming,
                ("label", "reference", "product_name", "origin", "tracking_code"),
            )
        return existing

    def apply_tax(self, key, label, month, tracking, amount, priority, source):
        month_key = month or SENTINEL
        sales = [
            order
            for order in self.orders.values()
            if order["order_key"] == key and order.get("sale_amount") is not None
        ]
        same_month = [order for order in sales if order["finance_month_key"] == month_key]
        if same_month:
            target = same_month[0]
        elif sales:
            earlier = [
                order
                for order in sales
                if month and order.get("finance_month") and order["finance_month"] <= month
            ]
            target = max(earlier or sales, key=lambda order: order["finance_month_key"])
        else:
            target = None
        if target is None:
            target = blank_order(key, label, "encomenda", month)
            target["sourcing_status"] = "TAXADO"
            target["priority"] = priority
            self.orders[self.identity(target)] = target
            self.report["taxas_sem_venda"] += 1
        else:
            self.report["taxas_fundidas"] += 1
        if amount is not None:
            old = target.get("import_tax")
            if old is not None and abs(old - amount) > 0.009:
                self.conflict(key, "import_tax", old, amount, source)
                if priority >= target["priority"]:
                    target["import_tax"] = amount
            elif old is None or priority >= target["priority"]:
                target["import_tax"] = amount
        if tracking and (priority >= target["priority"] or not target.get("tracking_code")):
            target["tracking_code"] = tracking
        target["tax_paid"] = True

    def set_goal(self, month, amount, source):
        if not month or amount is None:
            return
        previous = self.goals.get(month)
        if previous is not None and abs(previous - amount) > 0.009:
            self.conflict(month, "meta", previous, amount, source)
        self.goals[month] = amount

    def consume_finance(self, rows, month, priority, source):
        flow = "encomenda"
        tax_mode = False
        for row in rows:
            reference = clean_text(cell(row, 2))
            section = fold(reference or "")
            order_cell = cell(row, 1)
            key, label = parse_key(order_cell)
            if not key and section in {"loja nova", "nscreditos", "ns creditos", "taxas"}:
                if section == "taxas":
                    tax_mode = True
                elif section == "loja nova":
                    flow = "loja_nova"
                    tax_mode = False
                else:
                    flow = "ns_creditos"
                    tax_mode = False
                continue
            if not key:
                status = clean_text(cell(row, 3))
                if status and fold(status) in {"comprado", "taxado", "vendido", "enviado"}:
                    self.report["linhas_sem_numero"] += 1
                    if len(self.report["linhas_sem_numero_amostra"]) < 8:
                        self.report["linhas_sem_numero_amostra"].append(
                            {
                                "origem": source,
                                "status": status,
                                "referencia": reference,
                            }
                        )
                continue
            self.report["linhas_lidas"] += 1
            status = clean_text(cell(row, 3))
            sale = parse_money(cell(row, 4))
            purchase = parse_money(cell(row, 5))
            payment_fee = parse_money(cell(row, 6))
            shipping = parse_money(cell(row, 7))
            import_tax = parse_money(cell(row, 8))
            tracking = reference if reference and TRACKING_RE.match(reference) else None
            if tax_mode or (status and fold(status) == "taxado" and sale is None):
                self.pending_taxes.append(
                    (
                        key,
                        label,
                        month,
                        tracking or (reference if tax_mode else None),
                        import_tax,
                        priority,
                        source,
                    )
                )
                continue
            order = blank_order(key, label, flow, month)
            order["reference"] = None if tracking else reference
            order["tracking_code"] = tracking
            order["product_name"] = hint_from_label(label, key)
            order["sourcing_status"] = status
            order["sale_amount"] = sale
            order["purchase_amount"] = purchase
            order["payment_fee"] = payment_fee
            order["shipping_cost"] = shipping
            order["import_tax"] = import_tax
            if import_tax is not None:
                order["tax_paid"] = True
            if status and fold(status) == "comprado":
                order["purchased"] = True
            self.upsert_sale(order, priority, source)

    def pick(self, key, create=True):
        matches = [order for order in self.orders.values() if order["order_key"] == key]
        pool = [order for order in matches if order["flow"] == "encomenda"] or matches
        if not pool:
            if not create:
                return None
            order = blank_order(key, key, "encomenda", None)
            self.orders[self.identity(order)] = order
            return order
        return max(pool, key=lambda order: order["finance_month_key"])

    def patch(self, key, label, updates, force_text=()):
        if not key:
            meaningful = bool(label) or any(
                value not in (None, "", False)
                for field, value in updates.items()
                if field not in {"purchased", "cpf_linked", "tax_paid", "delivered"}
            )
            if meaningful:
                self.report["linhas_sem_numero"] += 1
            return None
        self.report["linhas_lidas"] += 1
        order = self.pick(key, create=True)
        if label and (not order.get("label") or order["label"] == order["order_key"]):
            order["label"] = label
        elif label and label not in {order.get("label"), order["order_key"]}:
            if not str(order.get("label") or "").lower().startswith("pedido"):
                order["label"] = label
        for field in ("purchased", "cpf_linked", "tax_paid", "delivered"):
            if updates.get(field):
                order[field] = True
        for field, value in updates.items():
            if field in {"purchased", "cpf_linked", "tax_paid", "delivered"}:
                continue
            if value in (None, ""):
                continue
            if field in force_text or not order.get(field):
                order[field] = value
        if order["delivered"] and not order.get("commercial_status"):
            order["commercial_status"] = "ENTREGUE"
        return order

    def read_finance(self, workbook):
        year = workbook["Financeiro 2026"]
        rows = [list(row) for row in year.iter_rows(values_only=True)]
        for month, block in split_blocks(rows):
            self.set_goal(month, find_meta(block), "Financeiro 2026")
            self.consume_finance(block, month, 1, "Financeiro 2026")
        for name, month in SHEET_MONTH.items():
            sheet = workbook[name]
            block = [list(row) for row in sheet.iter_rows(values_only=True)]
            self.set_goal(month, find_meta(block), name)
            self.consume_finance(block, month, 2, name)
        for tax in self.pending_taxes:
            self.apply_tax(*tax)

    def read_entregues(self, workbook):
        rows = list(workbook["Entregues"].iter_rows(values_only=True))
        for row in rows[1:]:
            key, label = parse_key(cell(row, 0))
            tracking = clean_text(cell(row, 4))
            shipment = clean_text(cell(row, 3))
            if shipment and TRACKING_RE.match(shipment) and not tracking:
                tracking = shipment
            self.patch(
                key,
                label,
                {
                    "origin": clean_text(cell(row, 1)),
                    "purchased": parse_bool(cell(row, 2)),
                    "tracking_code": tracking,
                    "cpf_linked": parse_bool(cell(row, 5)),
                    "tax_paid": parse_bool(cell(row, 6)),
                    "delivered": parse_bool(cell(row, 7)) is not False,
                    "notes_human": clean_text(cell(row, 8)),
                    "product_name": hint_from_label(label, key) if key else None,
                },
            )

    def read_pedidos(self, workbook):
        rows = list(workbook["Pedidos"].iter_rows(values_only=True))
        for row in rows[1:]:
            key, label = parse_key(cell(row, 0))
            product = clean_text(cell(row, 4))
            tracking = product if product and TRACKING_RE.match(product) else None
            self.patch(
                key,
                label,
                {
                    "origin": clean_text(cell(row, 1)),
                    "purchased": parse_bool(cell(row, 2)),
                    "supplier_ref": clean_text(cell(row, 3)),
                    "tracking_code": tracking,
                    "product_name": None if tracking else product,
                    "cpf_linked": parse_bool(cell(row, 5)),
                    "tax_paid": parse_bool(cell(row, 6)),
                    "delivered": parse_bool(cell(row, 7)),
                    "notes_robot": clean_text(cell(row, 8)),
                    "notes_human": clean_text(cell(row, 9)),
                },
            )

    def read_acompanhamentos(self, workbook):
        rows = list(workbook["Acompanhamentos"].iter_rows(values_only=True))
        for row in rows[1:]:
            key, label = parse_key(cell(row, 0))
            status = clean_text(cell(row, 3))
            if status:
                status = status.strip().upper()
            days = cell(row, 7)
            supplier_days = int(days) if isinstance(days, (int, float)) and not isinstance(days, bool) else None
            self.patch(
                key,
                label,
                {
                    "purchase_date": parse_date(cell(row, 1)),
                    "payment_date": parse_date(cell(row, 2)),
                    "commercial_status": status,
                    "product_name": clean_text(cell(row, 4)),
                    "notes_internal": clean_text(cell(row, 5)),
                    "supplier_days": supplier_days,
                    "origin": clean_text(cell(row, 8)),
                },
                force_text=("commercial_status", "product_name", "origin", "notes_internal"),
            )

    def read_rastreamento(self, workbook):
        rows = list(workbook["Rastreamento"].iter_rows(values_only=True))
        for row in rows[1:]:
            key, label = parse_key(cell(row, 1))
            event = None
            for value in row[7:]:
                event = parse_datetime(value)
                if event:
                    break
            self.patch(
                key,
                label,
                {
                    "origin": clean_text(cell(row, 2)),
                    "tracking_code": clean_text(cell(row, 3)),
                    "tracking_correios": clean_text(cell(row, 4)),
                    "tracking_situation": clean_text(cell(row, 5)),
                    "tracking_alert": clean_text(cell(row, 6)),
                    "tracking_event_at": event,
                },
                force_text=(
                    "tracking_code",
                    "tracking_correios",
                    "tracking_situation",
                    "tracking_alert",
                    "tracking_event_at",
                    "origin",
                ),
            )

    def read_notas(self, workbook):
        name = next(sheet for sheet in workbook.sheetnames if sheet.lower().startswith("notas de intermedia"))
        rows = list(workbook[name].iter_rows(values_only=True))
        for row in rows[1:]:
            key, label = parse_key(cell(row, 0))
            if not key:
                if any(cell(row, index) not in (None, "") for index in range(0, 4)):
                    self.report["linhas_sem_numero"] += 1
                continue
            self.report["linhas_lidas"] += 1
            order = self.pick(key, create=True)
            if label and not order.get("label"):
                order["label"] = label
            purchased = parse_bool(cell(row, 2))
            if purchased:
                order["purchased"] = True
            if not order.get("origin"):
                order["origin"] = clean_text(cell(row, 1))
            tracking = clean_text(cell(row, 3))
            if tracking and not order.get("tracking_code"):
                order["tracking_code"] = tracking
            for field, index in (
                ("purchase_amount", 4),
                ("payment_fee", 5),
                ("shipping_cost", 6),
            ):
                amount = parse_money(cell(row, index))
                if amount is not None and order.get(field) is None:
                    order[field] = amount

    def find_order_id(self, key):
        if not key:
            return None
        matches = [order for order in self.orders.values() if order["order_key"] == key]
        if not matches:
            return None
        return max(matches, key=lambda order: order["finance_month_key"])["id"]

    def read_cancelamentos(self, workbook):
        rows = list(workbook["Cancelamentos"].iter_rows(values_only=True))
        for row in rows:
            key, _label = parse_key(cell(row, 0))
            if not key:
                continue
            self.cancellations.append(
                {
                    "id": str(uuid.uuid4()),
                    "order_id": self.find_order_id(key),
                    "order_key": key,
                    "model": clean_text(cell(row, 1)),
                    "refund_method": clean_text(cell(row, 2)),
                    "bank_details": clean_text(cell(row, 3)),
                    "amount": parse_money(cell(row, 4)),
                    "due_date": parse_date(cell(row, 5)),
                    "done_date": parse_date(cell(row, 6)),
                    "reason": clean_text(cell(row, 7)),
                    "status": clean_text(cell(row, 8)),
                    "window_note": clean_text(cell(row, 9)),
                    "gateway": clean_text(cell(row, 10)),
                }
            )

    def read_estoque(self, workbook):
        rows = list(workbook["Estoque"].iter_rows(values_only=True))
        for row in rows:
            model = clean_text(cell(row, 3))
            brand = clean_text(cell(row, 2))
            if not model or fold(model) == "modelo":
                continue
            if fold(brand or "") in {"marca", "new store estoque"}:
                continue
            sale_raw = cell(row, 6)
            cost_raw = cell(row, 7)
            tracking = clean_text(cell(row, 5))
            order_id = None
            if tracking:
                for order in self.orders.values():
                    if order.get("tracking_code") and fold(order["tracking_code"]) == fold(tracking):
                        order_id = order["id"]
                        break
            self.inventory.append(
                {
                    "id": str(uuid.uuid4()),
                    "purchased_at": parse_date(cell(row, 1)),
                    "brand": brand,
                    "model": model,
                    "origin": clean_text(cell(row, 4)),
                    "tracking_code": tracking,
                    "sale_price": parse_money(sale_raw),
                    "sale_price_note": None if parse_money(sale_raw) is not None else clean_text(sale_raw),
                    "cost": parse_money(cost_raw),
                    "cost_note": None if parse_money(cost_raw) is not None else clean_text(cost_raw),
                    "location": clean_text(cell(row, 8)),
                    "notes": clean_text(cell(row, 9)),
                    "status": clean_text(cell(row, 10)),
                    "order_id": order_id,
                }
            )

    def read_trade_ins(self, workbook):
        name = next(sheet for sheet in workbook.sheetnames if "comprados" in fold(sheet))
        rows = list(workbook[name].iter_rows(values_only=True))
        inventory_by_code = {}
        for row in rows[1:]:
            client = clean_text(cell(row, 0))
            if not client or fold(client) == "cliente":
                continue
            code = clean_text(cell(row, 3))
            key, _label = parse_key(code) if code else (None, None)
            item = {
                "id": str(uuid.uuid4()),
                "client_name": client,
                "phone": clean_text(cell(row, 1)),
                "cpf": clean_text(cell(row, 2)),
                "code": key or code,
                "model": clean_text(cell(row, 4)),
                "condition": clean_text(cell(row, 5)),
                "address": clean_text(cell(row, 6)),
                "cost": parse_money(cell(row, 7)),
                "invoice_received": parse_bool(cell(row, 8)) is True,
                "delivered": parse_bool(cell(row, 9)) is True,
                "order_id": self.find_order_id(key) if key else None,
                "inventory_item_id": None,
            }
            self.trade_ins.append(item)
            if item["code"]:
                inventory_by_code[item["code"]] = item["id"]

    def finalize(self):
        for order in self.orders.values():
            if order["delivered"] and not order.get("commercial_status"):
                order["commercial_status"] = "ENTREGUE"
            order.pop("priority", None)
        flows = Counter(order["flow"] for order in self.orders.values())
        months = Counter(
            order["finance_month"] or "sem_mes" for order in self.orders.values()
        )
        self.report["pedidos_unicos"] = len(self.orders)
        self.report["por_fluxo"] = dict(flows)
        self.report["meses"] = dict(months)
        self.report["metas"] = self.goals
        self.report["estoque"] = len(self.inventory)
        self.report["compras_clientes"] = len(self.trade_ins)
        self.report["cancelamentos"] = len(self.cancellations)

    def payload(self):
        return {
            "orders": list(self.orders.values()),
            "goals": [
                {"month": month, "target_amount": amount}
                for month, amount in sorted(self.goals.items())
            ],
            "inventory": self.inventory,
            "trade_ins": self.trade_ins,
            "cancellations": self.cancellations,
        }


def build(path: Path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    importer = Importer()
    importer.read_finance(workbook)
    importer.read_entregues(workbook)
    importer.read_pedidos(workbook)
    importer.read_acompanhamentos(workbook)
    importer.read_rastreamento(workbook)
    importer.read_notas(workbook)
    importer.read_cancelamentos(workbook)
    importer.read_estoque(workbook)
    importer.read_trade_ins(workbook)
    importer.finalize()
    return importer.payload(), importer.report


def chunks(items, size):
    for index in range(0, len(items), size):
        yield items[index : index + size]


def load(payload):
    from supabase import create_client

    url = os.environ["SUPABASE_URL"]
    key = os.environ["SUPABASE_ANON_KEY"]
    email = os.environ["CONTROLE_EMAIL"]
    password = os.environ["CONTROLE_PASSWORD"]
    client = create_client(url, key)
    auth = client.auth.sign_in_with_password({"email": email, "password": password})
    if not auth.session:
        raise SystemExit("login do time falhou")
    client.table("ctl_cancellations").delete().neq("id", "00000000-0000-0000-0000-000000000000").execute()
    client.table("ctl_trade_ins").delete().neq("id", "00000000-0000-0000-0000-000000000000").execute()
    client.table("ctl_inventory_items").delete().neq("id", "00000000-0000-0000-0000-000000000000").execute()
    client.table("ctl_orders").delete().neq("id", "00000000-0000-0000-0000-000000000000").execute()
    client.table("ctl_monthly_goals").delete().neq("month", "1900-01-01").execute()
    for batch in chunks(payload["goals"], 100):
        client.table("ctl_monthly_goals").insert(batch).execute()
    for batch in chunks(payload["orders"], 200):
        client.table("ctl_orders").insert(batch).execute()
    for batch in chunks(payload["inventory"], 200):
        client.table("ctl_inventory_items").insert(batch).execute()
    for batch in chunks(payload["trade_ins"], 200):
        client.table("ctl_trade_ins").insert(batch).execute()
    for batch in chunks(payload["cancellations"], 200):
        client.table("ctl_cancellations").insert(batch).execute()


def load_env_file():
    env_path = ROOT / ".env.local"
    if not env_path.exists():
        return
    for line in env_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip())


def main():
    load_env_file()
    path = Path(os.environ.get("NS_XLSX", DEFAULT_XLSX))
    payload, report = build(path)
    data_dir = ROOT / "data"
    data_dir.mkdir(exist_ok=True)
    (data_dir / "import_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (data_dir / "import_payload.json").write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )
    printable = dict(report)
    print(json.dumps(printable, ensure_ascii=False, indent=2))
    if os.environ.get("SUPABASE_URL"):
        load(payload)
        print("carga concluida")


if __name__ == "__main__":
    main()
